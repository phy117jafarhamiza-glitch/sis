import io
import re
import difflib

import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, column_index_from_string

st.set_page_config(page_title="نقل الدرجات: الماستر شيت ← SIS", page_icon="📊", layout="wide")

st.markdown(
    """
    <style>
    html, body, [class*="css"] { direction: rtl; text-align: right; }
    .stDataFrame, .stDataEditor { direction: ltr; }
    </style>
    """,
    unsafe_allow_html=True,
)

NONE_OPT = "— لا شيء —"
SPECIAL_KEYS = ("نتيجه", "ملاحظات", "معدل", "ترتيب")

# ----------------------------------------------------------------- أدوات مساعدة
_TASHKEEL = re.compile(r"[ً-ْـٰ]")
_AR_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")


def norm(s) -> str:
    """توحيد النص العربي للمقارنة."""
    if s is None:
        return ""
    s = str(s).translate(_AR_DIGITS)
    s = _TASHKEEL.sub("", s)
    s = re.sub("[أإآٱ]", "ا", s)
    s = s.replace("ى", "ي").replace("ة", "ه").replace("ؤ", "و").replace("ئ", "ي")
    s = re.sub(r"[^\w\s]", " ", s)
    s = re.sub(r"\s+", " ", s).strip().lower()
    return s


def norm_label(s) -> str:
    s = norm(s)
    words = [w[2:] if w.startswith("ال") and len(w) > 3 else w for w in s.split()]
    return " ".join(words)


def open_wb(data: bytes, password: str, data_only: bool):
    bio = io.BytesIO(data)
    try:
        import msoffcrypto

        office = msoffcrypto.OfficeFile(bio)
        if office.is_encrypted():
            if not password:
                raise ValueError("الملف محمي بكلمة سر للفتح، أدخل كلمة السر.")
            out = io.BytesIO()
            office.load_key(password=password)
            office.decrypt(out)
            bio = out
    except ImportError:
        pass
    bio.seek(0)
    return openpyxl.load_workbook(bio, data_only=data_only)


def merged_map(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        anchor = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = anchor
    return m


def parse_rows(txt, default):
    rows = []
    for part in re.split(r"[,\s،]+", (txt or "").strip()):
        if not part:
            continue
        try:
            if "-" in part:
                a, b = part.split("-", 1)
                rows += list(range(int(a), int(b) + 1))
            else:
                rows.append(int(part))
        except ValueError:
            return default
    return sorted(set(rows)) or default


def build_labels(ws, rows):
    mm = merged_map(ws)
    labels = []
    for c in range(1, ws.max_column + 1):
        parts = []
        for r in rows:
            v = mm[(r, c)] if (r, c) in mm else ws.cell(r, c).value
            if v not in (None, ""):
                s = re.sub(r"\s+", " ", str(v)).strip()
                if not parts or parts[-1] != s:
                    parts.append(s)
        labels.append(" | ".join(parts))
    return labels


def col_options(labels):
    return [f"{get_column_letter(i + 1)}: {l}" if l else f"{get_column_letter(i + 1)}: (بدون عنوان)"
            for i, l in enumerate(labels)]


def letter_of(opt):
    return opt.split(":")[0].strip()


def col_index(letter):
    return column_index_from_string(letter)


def detect_layout(ws):
    """يحاول اكتشاف صف عناوين الأعمدة وصف أسماء المقررات تلقائياً."""
    sr = cr = None
    for r in range(1, min(ws.max_row, 40) + 1):
        for c in range(1, ws.max_column + 1):
            t = norm(ws.cell(r, c).value)
            if t == "اسم الطالب" and sr is None:
                sr = r
            if t == "اسم المقرر الدراسي" and cr is None:
                cr = r
    if sr is None:
        return [1], 2
    rows = [cr - 1, cr, sr] if cr and cr > 1 else [sr]
    return sorted(set(rows)), sr + 1


def find_col(labels, *keys):
    for i, l in enumerate(labels):
        n = norm(l.split(" | ")[-1])
        if any(n == k for k in keys):
            return i
    return None


def raw_preview(ws, n=14):
    rows = []
    for r in range(1, min(ws.max_row, n) + 1):
        rows.append([ws.cell(r, c).value for c in range(1, min(ws.max_column, 40) + 1)])
    df = pd.DataFrame(rows, index=range(1, len(rows) + 1))
    df.columns = [get_column_letter(i + 1) for i in range(df.shape[1])]
    return df.astype(str).replace("None", "")


# ------------------------------------------------- الطلبة وكتل المحاولات
def student_blocks(ws, first_row, name_col, att_col):
    """يرجع قائمة (الاسم، {رقم المحاولة: رقم الصف}) لكل طالب."""
    res, r, maxr = [], first_row, ws.max_row
    while r <= maxr:
        nm = ws.cell(r, name_col).value
        if nm in (None, ""):
            r += 1
            continue
        a = 1
        if att_col:
            try:
                a = int(ws.cell(r, att_col).value)
            except (TypeError, ValueError):
                a = 1
        rows = {a: r}
        k = r + 1
        while att_col and k <= maxr and ws.cell(k, name_col).value in (None, ""):
            try:
                av = int(ws.cell(k, att_col).value)
            except (TypeError, ValueError):
                break
            if av in rows:
                break
            rows[av] = k
            k += 1
        res.append((nm, rows))
        r = k
    return res


def name_key(s):
    return norm(s).replace(" ", "")


def name_score(a, b):
    a, b = name_key(a), name_key(b)
    if not a or not b:
        return 0.0
    short, long_ = (a, b) if len(a) <= len(b) else (b, a)
    full = difflib.SequenceMatcher(None, a, b).ratio()
    pref = difflib.SequenceMatcher(None, short, long_[: len(short) + 1]).ratio()
    return max(full, pref)


def match_students(tgt_blocks, src_blocks, threshold):
    """مطابقة شاملة: أعلى تشابه أولاً، كل طالب يُستخدم مرة واحدة."""
    cand = []
    for ti, (tn, _) in enumerate(tgt_blocks):
        for si, (sn, _) in enumerate(src_blocks):
            sc = name_score(tn, sn)
            if sc >= threshold:
                cand.append((sc, ti, si))
    cand.sort(reverse=True)
    t_used, s_used, out = set(), set(), {}
    for sc, ti, si in cand:
        if ti in t_used or si in s_used:
            continue
        t_used.add(ti)
        s_used.add(si)
        out[ti] = (si, sc)
    return out


# ------------------------------------------------------ اقتراح ربط الأعمدة
def col_kind(label):
    parts = label.split(" | ") if label else []
    if not parts:
        return None, None
    last = norm(parts[-1])
    if len(parts) <= 2:
        for k in SPECIAL_KEYS:
            if k in last and last != "الدرجه":
                return "special", k
    if last == "الدرجه" and len(parts) >= 2:
        return "subject", None
    return None, None


def subject_groups(labels, skip):
    """أعمدة المقررات مرتبة مع مجموعة الفصل (يملأ الفصل الناقص من اليسار)."""
    groups, cur = [], ""
    for i, l in enumerate(labels):
        if (i + 1) in skip:
            continue
        kind, _ = col_kind(l)
        if kind != "subject":
            continue
        parts = l.split(" | ")
        if len(parts) >= 3:
            cur = norm_label(parts[0])
        groups.append((i, cur))
    return groups


def suggest_mapping(t_labels, t_opts, s_labels, s_opts, skip_t, skip_s, mode):
    result = {}
    # أعمدة (المعدل/النتيجة/الملاحظات/الترتيب) تُربط دائماً بالاسم
    s_special = {}
    for i, l in enumerate(s_labels):
        if (i + 1) in skip_s:
            continue
        kind, key = col_kind(l)
        if kind == "special":
            s_special.setdefault(key, s_opts[i])
    for i, l in enumerate(t_labels):
        if (i + 1) in skip_t:
            continue
        kind, key = col_kind(l)
        if kind == "special" and key in s_special:
            result[t_opts[i]] = s_special[key]

    if mode == "بالترتيب داخل كل فصل":
        tg, sg = subject_groups(t_labels, skip_t), subject_groups(s_labels, skip_s)
        names = []
        for _, g in tg:
            if g not in names:
                names.append(g)
        s_names = []
        for _, g in sg:
            if g not in s_names:
                s_names.append(g)
        for k, g in enumerate(names):
            sg_key = s_names[k] if k < len(s_names) else None
            t_cols = [i for i, gg in tg if gg == g]
            s_cols = [i for i, gg in sg if gg == sg_key] if sg_key is not None else []
            for ti, si in zip(t_cols, s_cols):
                result[t_opts[ti]] = s_opts[si]
    else:
        src_norm = {}
        for i, l in enumerate(s_labels):
            if (i + 1) in skip_s or not l:
                continue
            src_norm.setdefault(norm_label(l), s_opts[i])
        for i, l in enumerate(t_labels):
            if (i + 1) in skip_t or not l or t_opts[i] in result:
                continue
            n = norm_label(l)
            if n in src_norm:
                result[t_opts[i]] = src_norm[n]
            else:
                best = difflib.get_close_matches(n, list(src_norm.keys()), n=1, cutoff=0.85)
                if best:
                    result[t_opts[i]] = src_norm[best[0]]
    return result


# ---------------------------------------------------------------- النقل
def transfer(src_ws, tgt_bytes, tgt_pwd, tgt_sheet, src_blocks, tgt_blocks, matches,
             col_pairs, avg_cols, overwrite, skip_empty, round_digits):
    out_wb = open_wb(tgt_bytes, tgt_pwd, False)
    out_ws = out_wb[tgt_sheet]
    written = skipped_merged = 0
    for ti, (si, _) in matches.items():
        _, t_rows = tgt_blocks[ti]
        _, s_rows = src_blocks[si]
        for att, s_row in s_rows.items():
            t_row = t_rows.get(att)
            if t_row is None:
                continue
            for t_col, s_col in col_pairs:
                v = src_ws.cell(s_row, s_col).value
                if skip_empty and v in (None, ""):
                    continue
                if round_digits is not None and t_col in avg_cols and isinstance(v, float):
                    v = round(v, round_digits)
                cell = out_ws.cell(t_row, t_col)
                if isinstance(cell, MergedCell):
                    skipped_merged += 1
                    continue
                if not overwrite and cell.value not in (None, ""):
                    continue
                cell.value = v
                written += 1
    buf = io.BytesIO()
    out_wb.save(buf)
    return buf.getvalue(), written, skipped_merged


# ------------------------------------------------------------------- الواجهة
def is_sis_file(data, pwd):
    """ملف SIS يحتوي أعمدة Course ID / Student Academic Status."""
    try:
        wb = open_wb(data, pwd, True)
        ws = wb.worksheets[0]
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 14), values_only=True):
            for v in row:
                if isinstance(v, str) and ("course id" in v.lower() or "student academic status" in v.lower()):
                    return True
    except Exception:
        pass
    return False


def file_panel(title, key, up, pwd, need_data_only):
    st.subheader(title)
    st.caption(f"الملف: {up.name}")
    data = up.getvalue()
    try:
        wb = open_wb(data, pwd, need_data_only)
    except Exception as e:
        st.error(f"تعذر فتح الملف: {e}")
        return None
    sheet = st.selectbox("الورقة (Sheet)", wb.sheetnames, key=f"sh_{key}")
    ws = wb[sheet]
    with st.expander("معاينة أول الصفوف"):
        st.dataframe(raw_preview(ws), use_container_width=True)
    d_rows, d_first = detect_layout(ws)
    rows_txt = st.text_input(
        "صفوف العناوين (مثال: 2,3,7 أو 4-7) — اكتُشفت تلقائياً",
        value=",".join(map(str, d_rows)), key=f"hr_{key}_{sheet}",
    )
    h_rows = parse_rows(rows_txt, d_rows)
    first_row = int(st.number_input("أول صف بيانات", 1, max(ws.max_row, 1), d_first, key=f"fr_{key}_{sheet}"))
    labels = build_labels(ws, h_rows)
    opts = col_options(labels)
    n_i = find_col(labels, "اسم الطالب")
    a_i = find_col(labels, "المحاوله")
    id_i = find_col(labels, "رقم الطالب")
    name_opt = st.selectbox("عمود اسم الطالب", opts, index=n_i or 0, key=f"nm_{key}_{sheet}")
    att_opt = st.selectbox("عمود رقم المحاولة (دور أول/ثاني/ثالث)", [NONE_OPT] + opts,
                           index=(a_i + 1) if a_i is not None else 0, key=f"at_{key}_{sheet}")
    return dict(
        name=up.name, data=data, pwd=pwd, wb=wb, ws=ws, sheet=sheet, first_row=first_row, rows=h_rows,
        labels=labels, opts=opts, name_col=col_index(letter_of(name_opt)),
        att_col=None if att_opt == NONE_OPT else col_index(letter_of(att_opt)),
        id_col=(id_i + 1) if id_i is not None else None,
    )


if __name__ == "__main__" or True:
    st.title("📊 نقل الدرجات من ماستر شيت اللجنة الامتحانية إلى ماستر شيت SIS")
    st.caption(
        "يطابق البرنامج الطلبة بالاسم (حتى لو كان الاسم مختصراً أو بإملاء مختلف قليلاً) ثم ينقل "
        "درجات الأدوار الأول والثاني والثالث والمعدل والنتيجة والملاحظات كما هي، دون تعديل أي خلية أخرى."
    )

    st.markdown("#### ارفع الملفين (بأي ترتيب) — يتعرف البرنامج تلقائياً على ملف SIS")
    u1, u2 = st.columns(2)
    f1 = u1.file_uploader("الملف الأول (xlsx)", type=["xlsx", "xlsm"], key="up_a")
    p1 = u1.text_input("كلمة سر الفتح (إن وجدت)", type="password", key="pw_a")
    f2 = u2.file_uploader("الملف الثاني (xlsx)", type=["xlsx", "xlsm"], key="up_b")
    p2 = u2.text_input("كلمة سر الفتح (إن وجدت)", type="password", key="pw_b")
    if not (f1 and f2):
        st.info("ارفع الملفين للمتابعة.")
        st.stop()

    s1, s2 = is_sis_file(f1.getvalue(), p1), is_sis_file(f2.getvalue(), p2)
    if s1 and not s2:
        (f_src, p_src), (f_tgt, p_tgt) = (f2, p2), (f1, p1)
    elif s2 and not s1:
        (f_src, p_src), (f_tgt, p_tgt) = (f1, p1), (f2, p2)
    else:
        (f_src, p_src), (f_tgt, p_tgt) = (f1, p1), (f2, p2)
        st.warning("لم أستطع تحديد ملف SIS تلقائياً، افترضت أن الأول ملف اللجنة والثاني ملف SIS.")
    st.success(f"ملف اللجنة (مصدر الدرجات): {f_src.name}   |   ملف SIS (يُعدَّل ويُنزَّل): {f_tgt.name}")

    left, right = st.columns(2)
    with left:
        src = file_panel("1) ملف الدرجات (ماستر شيت اللجنة الامتحانية)", "src", f_src, p_src, True)
    with right:
        tgt = file_panel("2) ملف SIS الذي سيستلم الدرجات", "tgt", f_tgt, p_tgt, False)

    if not (src and tgt):
        st.stop()

    skip_t = {tgt["name_col"]} | ({tgt["att_col"]} if tgt["att_col"] else set()) | ({tgt["id_col"]} if tgt["id_col"] else set())
    skip_s = {src["name_col"]} | ({src["att_col"]} if src["att_col"] else set()) | ({src["id_col"]} if src["id_col"] else set())

    st.divider()
    st.subheader("3) ربط الأعمدة")
    mode = st.radio("طريقة اقتراح ربط المقررات", ["بالترتيب داخل كل فصل", "حسب تشابه الأسماء"], horizontal=True)
    suggest = suggest_mapping(tgt["labels"], tgt["opts"], src["labels"], src["opts"], skip_t, skip_s, mode)

    rows = []
    for i, l in enumerate(tgt["labels"]):
        if (i + 1) in skip_t or not l:
            continue
        kind, _ = col_kind(l)
        if kind is None:
            continue
        o = tgt["opts"][i]
        rows.append({"عمود SIS": o, "عمود الدرجات المقابل": suggest.get(o, NONE_OPT)})
    map_df = pd.DataFrame(rows)
    st.write(
        "الجدول يعرض أعمدة SIS (المقررات + النتيجة/المعدل/الملاحظات) وأمام كل عمود ما يقابله في ملف الدرجات. "
        "**راجع الربط جيداً** خاصة إذا اختلفت أسماء المقررات بين الملفين، وعدّله من القائمة."
    )
    src_choices = [NONE_OPT] + [o for i, o in enumerate(src["opts"]) if (i + 1) not in skip_s and src["labels"][i]]
    sig = f"{mode}{tgt['sheet']}{src['sheet']}{tgt['rows']}{src['rows']}{len(map_df)}"
    edited = st.data_editor(
        map_df, key=f"map_{sig}", use_container_width=True, hide_index=True,
        disabled=["عمود SIS"],
        column_config={"عمود الدرجات المقابل": st.column_config.SelectboxColumn(
            "عمود الدرجات المقابل", options=src_choices, required=True, width="large")},
    )

    st.subheader("4) خيارات")
    c1, c2, c3, c4 = st.columns(4)
    overwrite = c1.checkbox("استبدال القيم الموجودة في SIS", value=True)
    skip_empty = c2.checkbox("تجاهل الخلايا الفارغة في ملف الدرجات", value=True)
    thr = c3.slider("حد تشابه الأسماء", 0.70, 1.0, 0.85, 0.01)
    rd = c4.selectbox("تقريب المعدل (منازل عشرية)", ["بدون تقريب", "0", "1", "2", "3"])
    round_digits = None if rd == "بدون تقريب" else int(rd)

    if not st.button("🚀 نقل الدرجات", type="primary"):
        st.stop()

    src_blocks = student_blocks(src["ws"], src["first_row"], src["name_col"], src["att_col"])
    tgt_blocks = student_blocks(tgt["ws"], tgt["first_row"], tgt["name_col"], tgt["att_col"])
    matches = match_students(tgt_blocks, src_blocks, thr)

    col_pairs, avg_cols = [], set()
    for _, row in edited.iterrows():
        if row["عمود الدرجات المقابل"] != NONE_OPT:
            t = col_index(letter_of(row["عمود SIS"]))
            col_pairs.append((t, col_index(letter_of(row["عمود الدرجات المقابل"]))))
            if "معدل" in norm(row["عمود SIS"]):
                avg_cols.add(t)

    out_bytes, written, skipped_merged = transfer(
        src["ws"], tgt["data"], tgt["pwd"], tgt["sheet"], src_blocks, tgt_blocks, matches,
        col_pairs, avg_cols, overwrite, skip_empty, round_digits,
    )

    st.success(f"تمت مطابقة {len(matches)} طالب من {len(tgt_blocks)} في SIS، وكُتبت {written} قيمة.")
    st.download_button(
        "⬇️ تحميل ملف SIS (بنفس شكله مع الدرجات المنقولة)", data=out_bytes, file_name="SIS_" + tgt["name"],
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary",
    )

    review = []
    for ti, (si, sc) in sorted(matches.items()):
        review.append({"الاسم في SIS": tgt_blocks[ti][0], "الاسم في ملف الدرجات": src_blocks[si][0],
                       "التشابه": round(sc, 2)})
    with st.expander("نتيجة مطابقة الأسماء (راجعها)", expanded=True):
        st.dataframe(pd.DataFrame(review), use_container_width=True)
    un_t = [tgt_blocks[i][0] for i in range(len(tgt_blocks)) if i not in matches]
    if un_t:
        st.warning("طلبة في SIS لم يُعثر عليهم في ملف الدرجات: " + "، ".join(map(str, un_t)))
    used_s = {si for si, _ in matches.values()}
    un_s = [src_blocks[i][0] for i in range(len(src_blocks)) if i not in used_s]
    if un_s:
        with st.expander(f"طلبة في ملف الدرجات لم يُنقلوا ({len(un_s)})"):
            st.write("، ".join(map(str, un_s)))
